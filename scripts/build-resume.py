"""Build the editable resume from its public Markdown source.

Requires python-docx. Export the result to PDF and inspect the rendered page
before replacing public/AlexKourkoumelisResume.pdf; see README.md.
"""

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE


def inline(paragraph, text):
    """Support only the bold and link syntax used in docs/resume.md."""
    for token in re.split(r'(\*\*.*?\*\*|\[[^]]+\]\([^)]+\))', text):
        if token.startswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('['):
            match = re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)', token)
            label, url = match.groups()
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '234E83')
            underline = OxmlElement('w:u')
            underline.set(qn('w:val'), 'single')
            properties.extend([color, underline])
            run.append(properties)
            value = OxmlElement('w:t')
            value.text = label
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(token)


def build(source, output):
    document = Document()
    # The bundled Word defaults can carry a decorative Title paragraph border.
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.5)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.font.color.rgb = RGBColor.from_string('171717')
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.03

    for name, size, before, after in [
        ('Title', 22, 0, 2), ('Heading 1', 11, 8, 3), ('Heading 2', 11, 5, 1),
    ]:
        style = document.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('000000')
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, size in [('Contact', 10), ('Role', 10.5)]:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.size = Pt(size)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles['List Bullet']
    bullet.base_style = normal
    bullet.paragraph_format.left_indent = Inches(0.13)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.keep_together = True

    in_header = True
    current_section = ''
    for line in source.read_text().splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            paragraph = document.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            in_header = False
            current_section = line[3:]
            paragraph = document.add_paragraph(current_section.upper(), 'Heading 1')
        elif line.startswith('### '):
            paragraph = document.add_paragraph(line[4:], 'Heading 2')
        else:
            is_bullet = line.startswith('- ')
            style = 'List Bullet' if is_bullet else (
                'Contact' if in_header else 'Role' if current_section == 'Experience' else 'Normal'
            )
            paragraph = document.add_paragraph(style=style)
            inline(paragraph, line[2:] if is_bullet else line)
    document.core_properties.title = 'Alex Kourkoumelis Resume'
    document.core_properties.subject = 'Lead Software Engineer'
    document.core_properties.author = 'Alex Kourkoumelis'
    document.core_properties.last_modified_by = 'Alex Kourkoumelis'
    document.core_properties.created = document.core_properties.modified = datetime.now(timezone.utc)
    document.core_properties.keywords = 'React, TypeScript, frontend architecture, software engineering'
    document.core_properties.comments = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f'Wrote {output}')


if __name__ == '__main__':
    root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=root / 'docs/resume.md')
    parser.add_argument('--output', type=Path, default=root / 'docs/AlexKourkoumelisResume.docx')
    args = parser.parse_args()
    build(args.source, args.output)
